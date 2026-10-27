"""Resolve chat file_url into Gemini-safe HumanMessage content.

Images stay multimodal. DOCX/PDF are not valid Gemini image URLs — we download
them and either extract text (DOCX) or send a data-URL with the correct MIME (PDF).
"""

from __future__ import annotations

import base64
import io
import logging
import re
from pathlib import PurePosixPath
from typing import Any
from urllib.parse import unquote, urlparse

import httpx

logger = logging.getLogger(__name__)

_IMAGE_SUFFIXES = frozenset({".jpg", ".jpeg", ".png", ".gif", ".webp", ".img"})
_DOCX_SUFFIXES = frozenset({".docx"})
_PDF_SUFFIXES = frozenset({".pdf"})

_MAX_DOWNLOAD_BYTES = 8 * 1024 * 1024
_MAX_EXTRACT_CHARS = 80_000
_DOWNLOAD_TIMEOUT = 30.0

_MIME_BY_SUFFIX = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".img": "image/jpeg",
    ".pdf": "application/pdf",
    ".docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
}


def _suffix_from_url(url: str) -> str:
    path = unquote(urlparse(url).path or "")
    # Strip common signed-URL path noise; take last path segment.
    name = PurePosixPath(path).name
    if "." not in name:
        return ""
    return "." + name.rsplit(".", 1)[-1].lower()


def _guess_suffix(url: str, content_type: str | None) -> str:
    suffix = _suffix_from_url(url)
    if suffix in _IMAGE_SUFFIXES | _DOCX_SUFFIXES | _PDF_SUFFIXES:
        return suffix
    ct = (content_type or "").split(";")[0].strip().lower()
    if ct in {"image/jpeg", "image/jpg"}:
        return ".jpg"
    if ct == "image/png":
        return ".png"
    if ct == "image/gif":
        return ".gif"
    if ct == "image/webp":
        return ".webp"
    if ct == "application/pdf":
        return ".pdf"
    if "wordprocessingml" in ct or ct == "application/msword":
        return ".docx"
    return suffix


def _extract_docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _truncate(text: str, limit: int = _MAX_EXTRACT_CHARS) -> str:
    if len(text) <= limit:
        return text
    return text[: limit - 20].rstrip() + "\n\n[…truncated…]"


async def _download(url: str) -> tuple[bytes, str | None]:
    async with httpx.AsyncClient(follow_redirects=True, timeout=_DOWNLOAD_TIMEOUT) as client:
        resp = await client.get(url)
        resp.raise_for_status()
        data = resp.content
        if len(data) > _MAX_DOWNLOAD_BYTES:
            raise ValueError(f"Attachment too large ({len(data)} bytes)")
        return data, resp.headers.get("content-type")


def _download_sync(url: str) -> tuple[bytes, str | None]:
    with httpx.Client(follow_redirects=True, timeout=_DOWNLOAD_TIMEOUT) as client:
        resp = client.get(url)
        resp.raise_for_status()
        data = resp.content
        if len(data) > _MAX_DOWNLOAD_BYTES:
            raise ValueError(f"Attachment too large ({len(data)} bytes)")
        return data, resp.headers.get("content-type")


def _as_data_url(data: bytes, mime: str) -> str:
    b64 = base64.standard_b64encode(data).decode("ascii")
    return f"data:{mime};base64,{b64}"


def _text_with_document(message: str, filename: str, body: str) -> str:
    body = _truncate(body)
    header = f"[Attached document: {filename}]\n\n{body}" if body else f"[Attached document: {filename} — no extractable text]"
    msg = (message or "").strip()
    if not msg:
        return header
    return f"{msg}\n\n---\n{header}"


def _build_from_bytes(
    message: str,
    url: str,
    data: bytes,
    content_type: str | None,
) -> str | list[dict[str, Any]]:
    suffix = _guess_suffix(url, content_type)
    filename = PurePosixPath(unquote(urlparse(url).path or "")).name or f"attachment{suffix}"

    if suffix in _IMAGE_SUFFIXES:
        return [
            {"type": "text", "text": message or "Please review this image."},
            {"type": "image_url", "image_url": {"url": url, "detail": "high"}},
        ]

    if suffix in _DOCX_SUFFIXES:
        try:
            text = _extract_docx_text(data)
        except Exception:
            logger.exception("DOCX text extraction failed for %s", filename)
            text = ""
        if text:
            return _text_with_document(message, filename, text)
        # Last resort: data URL with DOCX MIME (some providers accept it).
        mime = _MIME_BY_SUFFIX[".docx"]
        return [
            {
                "type": "text",
                "text": message
                or "Please review the attached Word document under Ethiopian law.",
            },
            {
                "type": "image_url",
                "image_url": {"url": _as_data_url(data, mime), "detail": "high"},
            },
        ]

    if suffix in _PDF_SUFFIXES:
        mime = _MIME_BY_SUFFIX[".pdf"]
        return [
            {
                "type": "text",
                "text": message
                or "Please review the attached PDF under Ethiopian law.",
            },
            {
                "type": "image_url",
                "image_url": {"url": _as_data_url(data, mime), "detail": "high"},
            },
        ]

    # Unknown type: try to treat remote URL as image if it looks like one, else note failure.
    if re.search(r"\.(jpe?g|png|gif|webp)(\?|$)", url, re.I):
        return [
            {"type": "text", "text": message or "Please review this image."},
            {"type": "image_url", "image_url": {"url": url, "detail": "high"}},
        ]
    note = (
        f"{(message or '').strip()}\n\n"
        f"[Attached file `{filename}` could not be processed for review. "
        "Please paste the text or upload PDF/DOCX/JPG/PNG.]"
    ).strip()
    return note


async def resolve_human_message_content(
    message: str, file_url: str | None
) -> str | list[dict[str, Any]]:
    """Async resolver for legal-agent / chat streams."""
    url = (file_url or "").strip() or None
    if not url:
        return message

    suffix = _suffix_from_url(url)
    # Fast path: images can use the remote URL directly (no download).
    if suffix in _IMAGE_SUFFIXES:
        return [
            {"type": "text", "text": message or "Please review this image."},
            {"type": "image_url", "image_url": {"url": url, "detail": "high"}},
        ]

    try:
        data, content_type = await _download(url)
    except Exception as exc:
        logger.warning("Failed to download chat attachment %s: %s", url[:120], exc)
        base = (message or "").strip()
        err = (
            "[Could not download the attached file for analysis. "
            "Please try again or paste the document text.]"
        )
        return f"{base}\n\n{err}".strip() if base else err

    return _build_from_bytes(message, url, data, content_type)


def resolve_human_message_content_sync(
    message: str, file_url: str | None
) -> str | list[dict[str, Any]]:
    """Sync resolver for generation agents."""
    url = (file_url or "").strip() or None
    if not url:
        return message

    suffix = _suffix_from_url(url)
    if suffix in _IMAGE_SUFFIXES:
        return [
            {"type": "text", "text": message or "Please review this image."},
            {"type": "image_url", "image_url": {"url": url}},
        ]

    try:
        data, content_type = _download_sync(url)
    except Exception as exc:
        logger.warning("Failed to download chat attachment %s: %s", url[:120], exc)
        base = (message or "").strip()
        err = (
            "[Could not download the attached file for analysis. "
            "Please try again or paste the document text.]"
        )
        return f"{base}\n\n{err}".strip() if base else err

    return _build_from_bytes(message, url, data, content_type)
